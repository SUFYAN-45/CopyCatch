"""Document extraction service for TXT, PDF, and DOCX files."""

import io
from pathlib import Path
from typing import Union
import zipfile

import docx
import pypdf
from pypdf.errors import FileNotDecryptedError, PdfReadError


class DocumentExtractionError(Exception):
    """Base exception raised when document extraction fails."""


class DocumentEmptyError(DocumentExtractionError):
    """Raised when the document is empty or contains no extractable text."""


class UnsupportedFormatError(DocumentExtractionError):
    """Raised when an unsupported document format is encountered."""


class DocumentExtractor:
    """Extracts plain text from supported document formats (.txt, .pdf, .docx)."""

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}

    @classmethod
    def extract_text(cls, file_path: Union[str, Path]) -> str:
        """Extract all text from the specified document path."""
        path = Path(file_path)

        if not path.exists():
            raise DocumentExtractionError(f"File not found: '{path}'.")

        if path.stat().st_size == 0:
            raise DocumentEmptyError(f"Document '{path.name}' is empty (0 bytes).")

        ext = path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            allowed = ", ".join(sorted(cls.SUPPORTED_EXTENSIONS))
            raise UnsupportedFormatError(
                f"Unsupported format '{ext}'. Supported formats are: {allowed}."
            )

        if ext == ".txt":
            raw_text = cls.extract_txt(path)
        elif ext == ".pdf":
            raw_text = cls.extract_pdf(path)
        elif ext == ".docx":
            raw_text = cls.extract_docx(path)
        else:
            raise UnsupportedFormatError(f"Unhandled file extension: '{ext}'.")

        cleaned_text = raw_text.strip()
        if not cleaned_text:
            raise DocumentEmptyError(f"Document '{path.name}' contains no readable text content.")

        return cleaned_text

    @staticmethod
    def extract_txt(path: Path) -> str:
        """Extract text from a plain-text document with encoding fallback."""
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                return path.read_text(encoding="latin-1")
            except Exception:
                return path.read_text(encoding="utf-8", errors="replace")
        except Exception as err:
            raise DocumentExtractionError(f"Failed to read TXT file '{path.name}': {err}") from err

    @staticmethod
    def extract_pdf(path: Path) -> str:
        """Extract text from all pages of a PDF document using pypdf."""
        try:
            reader = pypdf.PdfReader(str(path))

            if reader.is_encrypted:
                try:
                    # Attempt decrypt with empty password for unprompted reading
                    decrypted = reader.decrypt("")
                    if decrypted == 0:
                        raise DocumentExtractionError("PDF is encrypted and password-protected.")
                except Exception as err:
                    raise DocumentExtractionError(f"PDF is encrypted: {err}") from err

            extracted_pages = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text.strip())

            return "\n\n".join(extracted_pages)

        except (PdfReadError, FileNotDecryptedError) as err:
            raise DocumentExtractionError(f"Corrupted or invalid PDF file '{path.name}': {err}") from err
        except Exception as err:
            raise DocumentExtractionError(f"Failed to extract PDF '{path.name}': {err}") from err

    @staticmethod
    def extract_docx(path: Path) -> str:
        """Extract text from paragraphs and tables of a DOCX document."""
        try:
            doc = docx.Document(str(path))
            elements = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    elements.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        elements.append(" | ".join(row_cells))

            return "\n\n".join(elements)

        except (zipfile.BadZipFile, KeyError) as err:
            raise DocumentExtractionError(f"Corrupted or invalid DOCX archive '{path.name}': {err}") from err
        except Exception as err:
            raise DocumentExtractionError(f"Failed to extract DOCX '{path.name}': {err}") from err


document_extractor = DocumentExtractor()
